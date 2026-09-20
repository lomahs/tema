"""Detecting a workbook's TOOL_DATA, writing it, and diffing it against what
is already there.

`parser.tool_data_builder` does the detecting; this covers the workbook-level
operations around it — reading the sheets, replacing the sheet in place, and
the comparison that turns "Create" into "Check" once a workbook has one.
"""

import pytest
from openpyxl import load_workbook

from tcm.infrastructure.excel.reader import TOOL_DATA_SHEET, parse_tool_data
from tcm.domain.case import SheetConfig
from tcm.infrastructure.excel.tool_data import detect_file, diff_configs, write_tool_data_sheet
from tcm.infrastructure.excel.workbook import has_tool_data, read_sheets
from tests.conftest import config_row, write_workbook


def sheet_cells():
    """A sheet laid out the way the real ones are: a device row above a header."""
    return {
        (2, "C"): "Pad(Flex)", (2, "H"): "27系Phone",
        (3, "A"): "No.", (3, "B"): "担当者",
        (3, "C"): "結果", (3, "D"): "確認日", (3, "E"): "確認者",
        (3, "F"): "チケットNo.", (3, "G"): "備考",
        (3, "H"): "結果", (3, "I"): "確認日", (3, "J"): "確認者",
        (3, "K"): "チケットNo.", (3, "L"): "備考",
        (5, "A"): "TC-1", (5, "B"): "FPT", (5, "C"): "OK",
        (6, "A"): "TC-2", (6, "B"): "FPT", (6, "C"): "NG",
    }


@pytest.fixture
def bare(tmp_path):
    """A workbook nobody has described yet — no TOOL_DATA sheet at all."""
    return write_workbook(tmp_path / "Bare.xlsx", None, {"Login": sheet_cells()})


def config(device, cols, **overrides):
    return SheetConfig(**config_row("Login", device, 5, 6, cols=cols, **overrides))


# --- reading ---------------------------------------------------------------

def test_has_tool_data_is_false_for_a_workbook_without_the_sheet(bare):
    assert has_tool_data(bare) is False


def test_has_tool_data_is_true_once_the_sheet_exists(tmp_path):
    path = write_workbook(
        tmp_path / "TC.xlsx",
        [config_row("Login", "iPhone", 4, 8)],
        {"Login": {(4, "A"): "TC-1"}},
    )

    assert has_tool_data(path) is True


def test_read_sheets_leaves_out_the_tool_data_sheet(tmp_path):
    path = write_workbook(
        tmp_path / "TC.xlsx",
        [config_row("Login", "iPhone", 4, 8)],
        {"Login": {(4, "A"): "TC-1"}},
    )

    assert list(read_sheets(path)) == ["Login"]


# --- detecting -------------------------------------------------------------

def test_detect_file_finds_both_device_blocks(bare):
    configs, unresolved = detect_file(bare)

    assert [(c.sheet, c.device) for c in configs] == [
        ("Login", "Pad(Flex)"), ("Login", "27系Phone"),
    ]
    assert unresolved == []


def test_detect_file_reads_each_block_column_letters(bare):
    configs, _ = detect_file(bare)

    pad = configs[0]
    assert (pad.test_no_col, pad.scope_col) == ("A", "B")
    assert (pad.result_col, pad.test_date_col, pad.pic_col) == ("C", "D", "E")
    assert (pad.ticket_id_col, pad.note_col) == ("F", "G")


def test_detect_file_spans_the_rows_below_the_header(bare):
    configs, _ = detect_file(bare)

    assert (configs[0].start_row, configs[0].end_row) == (5, 6)


def one_device_sheet(header_row, data_rows):
    """A one-device sheet with its data rows placed exactly where asked.

    The gap between a header row and the first case is not fixed in real
    workbooks — some sheets carry a sub-header row, some start immediately,
    some leave a spacer or two — so the tests below vary it.
    """
    cells = {
        (header_row - 1, "C"): "Pad(Flex)",
        (header_row, "A"): "No.", (header_row, "B"): "担当者",
        (header_row, "C"): "結果", (header_row, "D"): "確認日",
        (header_row, "E"): "確認者",
        (header_row, "F"): "チケットNo.", (header_row, "G"): "備考",
    }
    for row in data_rows:
        cells[(row, "A")] = f"TC-{row}"
        cells[(row, "B")] = "FPT"
        cells[(row, "C")] = "OK"
    return cells


def span(tmp_path, cells, name="TC.xlsx"):
    path = write_workbook(tmp_path / name, None, {"Login": cells})
    configs, _ = detect_file(path)
    return configs[0].start_row, configs[0].end_row


def test_detect_starts_on_the_first_case_when_it_sits_under_the_header(tmp_path):
    """No sub-header, no spacer: row 4 is a case and must not be skipped."""
    cells = one_device_sheet(header_row=3, data_rows=[4, 5, 6])

    assert span(tmp_path, cells) == (4, 6)


def test_detect_starts_on_the_first_case_below_a_wider_gap(tmp_path):
    """Two spacer rows, not one. The span must not open above the data."""
    cells = one_device_sheet(header_row=3, data_rows=[6, 7, 8])

    assert span(tmp_path, cells) == (6, 8)


def test_detect_steps_over_a_sub_header_row(tmp_path):
    """The row under the header repeats "No." rather than numbering a case."""
    cells = one_device_sheet(header_row=3, data_rows=[5, 6])
    cells[(4, "A")] = "No."

    assert span(tmp_path, cells) == (5, 6)


def test_a_sheet_with_no_cases_gets_a_span_that_is_not_inverted(tmp_path):
    """A header describing nothing yet still has to name a readable span."""
    cells = one_device_sheet(header_row=3, data_rows=[])

    start_row, end_row = span(tmp_path, cells)
    assert start_row <= end_row


# --- writing ---------------------------------------------------------------

def test_write_adds_a_tool_data_sheet_a_workbook_lacked(bare):
    configs, _ = detect_file(bare)

    write_tool_data_sheet(bare, configs)

    assert [(c.sheet, c.device) for c in parse_tool_data(bare)] == [
        ("Login", "Pad(Flex)"), ("Login", "27系Phone"),
    ]


def test_write_leaves_the_test_case_sheets_untouched(bare):
    configs, _ = detect_file(bare)

    write_tool_data_sheet(bare, configs)

    ws = load_workbook(bare)["Login"]
    assert ws["A5"].value == "TC-1"
    assert ws["C3"].value == "結果"


def test_write_replaces_an_existing_tool_data_sheet(tmp_path):
    path = write_workbook(
        tmp_path / "TC.xlsx",
        [config_row("Login", "stale", 99, 100)],
        {"Login": sheet_cells()},
    )
    configs, _ = detect_file(path)

    write_tool_data_sheet(path, configs)

    wb = load_workbook(path)
    assert wb.sheetnames.count(TOOL_DATA_SHEET) == 1
    assert [c.device for c in parse_tool_data(path)] == ["Pad(Flex)", "27系Phone"]


def test_write_can_target_a_different_path(bare, tmp_path):
    configs, _ = detect_file(bare)

    dest = write_tool_data_sheet(bare, configs, str(tmp_path / "out.xlsx"))

    assert dest == str(tmp_path / "out.xlsx")
    assert has_tool_data(dest) is True
    assert has_tool_data(bare) is False


# --- diffing ---------------------------------------------------------------

def test_identical_configs_diff_as_no_difference():
    existing = [config("Pad", "A B C D E F G")]
    detected = [config("Pad", "A B C D E F G")]

    diff = diff_configs(existing, detected)

    assert diff.differs is False
    assert diff.same == [("Login", "Pad")]


def test_a_changed_column_letter_is_reported_field_by_field():
    existing = [config("Pad", "A B C D E F G")]
    detected = [config("Pad", "A B C D E F H")]

    diff = diff_configs(existing, detected)

    assert diff.differs is True
    assert diff.changed == [{
        "sheet": "Login", "device": "Pad",
        "fields": [{"field": "note_col", "existing": "G", "detected": "H"}],
    }]


def test_a_changed_row_span_is_reported_too():
    existing = [config("Pad", "A B C D E F G")]
    detected = [SheetConfig(**config_row("Login", "Pad", 5, 40))]

    diff = diff_configs(existing, detected)

    assert diff.changed[0]["fields"] == [
        {"field": "end_row", "existing": 6, "detected": 40},
    ]


def test_a_block_only_the_sheet_knows_about_is_reported():
    existing = [config("Pad", "A B C D E F G"), config("Phone", "A B H I J K L")]
    detected = [config("Pad", "A B C D E F G")]

    diff = diff_configs(existing, detected)

    assert diff.only_existing == [("Login", "Phone")]
    assert diff.only_detected == []
    assert diff.differs is True


def test_a_newly_detected_block_is_reported():
    existing = [config("Pad", "A B C D E F G")]
    detected = [config("Pad", "A B C D E F G"), config("Phone", "A B H I J K L")]

    diff = diff_configs(existing, detected)

    assert diff.only_detected == [("Login", "Phone")]
    assert diff.only_existing == []


def test_diffing_a_workbook_against_its_own_detection_finds_nothing(bare):
    configs, _ = detect_file(bare)
    write_tool_data_sheet(bare, configs)

    diff = diff_configs(parse_tool_data(bare), detect_file(bare)[0])

    assert diff.differs is False
