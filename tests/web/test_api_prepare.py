"""The /api/prepare endpoints: listing the loaded workbooks, and the two
operations that write to them.

Both operations preview by default and write only when asked, so most of what
matters here is that a call without `apply` leaves the file on disk alone.
"""

import os

import pytest
from openpyxl import load_workbook

from tcm.web import routes
from app import create_app
from tcm.infrastructure.excel.reader import parse_tool_data
from tcm.infrastructure.excel.workbook import has_tool_data
from tests.conftest import config_row, write_workbook


@pytest.fixture
def client():
    app = create_app()
    app.config.update(TESTING=True)
    routes._data.update({"cases": [], "file_results": [], "source": None})
    with app.test_client() as c:
        yield c


def bare_cells():
    """A sheet with the human headers detection reads, and no TOOL_DATA."""
    return {
        (2, "C"): "Pad(Flex)",
        (3, "A"): "No.", (3, "B"): "担当者",
        (3, "C"): "結果", (3, "D"): "確認日", (3, "E"): "確認者",
        (3, "F"): "チケットNo.", (3, "G"): "備考",
        (5, "A"): "TC-1", (5, "B"): "FPT", (5, "C"): "OK",
        (6, "A"): "TC-2", (6, "B"): "FPT", (6, "C"): "対象外",
    }


@pytest.fixture
def folder(tmp_path):
    """One described workbook and one nobody has described yet."""
    write_workbook(
        tmp_path / "Described.xlsx",
        [config_row("Login", "iPhone", 4, 6, cols="A B C D E F G")],
        {"Login": {
            (4, "A"): "TC-1", (4, "B"): "FPT", (4, "C"): "OK",
            (4, "D"): "2026-08-05", (4, "E"): "lee",
            (5, "A"): "TC-2", (5, "B"): "FPT", (5, "C"): "対象外",
            (5, "D"): "2026-08-05", (5, "E"): "lee",
            (6, "A"): "TC-3", (6, "B"): "FPT", (6, "C"): "NG",
            (6, "D"): "2026-08-05", (6, "E"): "lee", (6, "F"): "BUG-1",
        }},
    )
    write_workbook(tmp_path / "Bare.xlsx", None, {"Login": bare_cells()})
    return str(tmp_path)


@pytest.fixture
def loaded(client, folder):
    """The folder loaded, the way the user would before opening the drawer."""
    client.post("/api/load", json={"folder": folder})
    return folder


def paths(client, name):
    """The server's own path for one file, as the UI would pass it back."""
    files = client.get("/api/prepare/files").get_json()["files"]
    return next(f["path"] for f in files if f["file"] == name)


# --- listing ---------------------------------------------------------------

def test_listing_before_anything_is_loaded_is_an_error(client):
    res = client.get("/api/prepare/files")

    assert res.status_code == 400
    assert "load" in res.get_json()["error"].lower()


def test_listing_names_every_workbook_in_the_loaded_folder(client, loaded):
    files = client.get("/api/prepare/files").get_json()["files"]

    assert sorted(f["file"] for f in files) == ["Bare.xlsx", "Described.xlsx"]


def test_listing_says_which_workbooks_carry_tool_data(client, loaded):
    files = {f["file"]: f for f in client.get("/api/prepare/files").get_json()["files"]}

    assert files["Described.xlsx"]["has_tool_data"] is True
    assert files["Bare.xlsx"]["has_tool_data"] is False


def test_listing_counts_the_device_blocks_a_workbook_describes(client, loaded):
    files = {f["file"]: f for f in client.get("/api/prepare/files").get_json()["files"]}

    assert files["Described.xlsx"]["blocks"] == 1
    assert files["Bare.xlsx"]["blocks"] is None


def test_listing_follows_a_source_loaded_as_an_explicit_file_list(client, folder):
    one = os.path.join(folder, "Described.xlsx")
    client.post("/api/load", json={"files": [one]})

    files = client.get("/api/prepare/files").get_json()["files"]

    assert [f["file"] for f in files] == ["Described.xlsx"]


def test_a_file_outside_an_explicit_file_list_source_is_refused(client, folder):
    client.post("/api/load", json={"files": [os.path.join(folder, "Described.xlsx")]})

    res = client.post(
        "/api/prepare/tool-data",
        json={"files": [os.path.join(folder, "Bare.xlsx")], "apply": True},
    )

    assert res.status_code == 400


def test_listing_notices_a_file_added_since_the_load(client, loaded, tmp_path):
    write_workbook(tmp_path / "Later.xlsx", None, {"Login": bare_cells()})

    files = client.get("/api/prepare/files").get_json()["files"]

    assert "Later.xlsx" in [f["file"] for f in files]


def test_listing_skips_excel_lock_files(client, tmp_path):
    write_workbook(tmp_path / "~$Described.xlsx", None, {"Login": bare_cells()})
    write_workbook(tmp_path / "Bare.xlsx", None, {"Login": bare_cells()})
    client.post("/api/load", json={"folder": str(tmp_path)})

    files = client.get("/api/prepare/files").get_json()["files"]

    assert [f["file"] for f in files] == ["Bare.xlsx"]


# --- creating TOOL_DATA ----------------------------------------------------

def test_tool_data_previews_without_writing(client, loaded):
    path = paths(client, "Bare.xlsx")

    res = client.post("/api/prepare/tool-data", json={"files": [path]})

    assert res.status_code == 200
    result = res.get_json()["results"][0]
    assert result["written"] is False
    assert [c["device"] for c in result["detected"]] == ["Pad(Flex)"]
    assert has_tool_data(path) is False


def test_tool_data_writes_the_sheet_when_applied(client, loaded):
    path = paths(client, "Bare.xlsx")

    res = client.post("/api/prepare/tool-data", json={"files": [path], "apply": True})

    assert res.get_json()["results"][0]["written"] is True
    assert [c.device for c in parse_tool_data(path)] == ["Pad(Flex)"]


def test_tool_data_reports_no_diff_for_a_workbook_that_has_none(client, loaded):
    path = paths(client, "Bare.xlsx")

    result = client.post("/api/prepare/tool-data", json={"files": [path]}).get_json()["results"][0]

    assert result["diff"] is None


def test_tool_data_diffs_against_a_sheet_already_in_the_workbook(client, loaded):
    path = paths(client, "Described.xlsx")

    result = client.post("/api/prepare/tool-data", json={"files": [path]}).get_json()["results"][0]

    # The described workbook has no human headers, so detection finds nothing
    # and the one block it already describes shows as dropped.
    assert result["diff"]["differs"] is True
    assert result["diff"]["only_existing"] == [{"sheet": "Login", "device": "iPhone"}]


def test_tool_data_refuses_to_write_when_nothing_was_detected(client, loaded):
    path = paths(client, "Described.xlsx")

    result = client.post(
        "/api/prepare/tool-data", json={"files": [path], "apply": True}
    ).get_json()["results"][0]

    assert result["written"] is False
    # The sheet it already had is still there, untouched.
    assert [c.device for c in parse_tool_data(path)] == ["iPhone"]


def test_tool_data_handles_every_file_in_one_call(client, loaded):
    files = [f["path"] for f in client.get("/api/prepare/files").get_json()["files"]]

    results = client.post("/api/prepare/tool-data", json={"files": files}).get_json()["results"]

    assert len(results) == 2


# --- clearing results ------------------------------------------------------

def test_clear_previews_without_writing(client, loaded):
    path = paths(client, "Described.xlsx")

    res = client.post("/api/prepare/clear", json={"files": [path], "keep": []})

    result = res.get_json()["results"][0]
    assert result["applied"] is False
    assert result["plan"]["rows"] == 3
    assert load_workbook(path)["Login"]["C4"].value == "OK"


def test_clear_keeps_cancel_when_no_keep_is_given(client, loaded):
    path = paths(client, "Described.xlsx")

    result = client.post(
        "/api/prepare/clear", json={"files": [path]}
    ).get_json()["results"][0]

    assert result["plan"]["rows"] == 2
    assert result["plan"]["kept"] == {"Cancel": 1}


def test_clear_blanks_the_result_cells_when_applied(client, loaded):
    path = paths(client, "Described.xlsx")

    client.post("/api/prepare/clear", json={"files": [path], "apply": True})

    ws = load_workbook(path)["Login"]
    assert [ws[f"{c}4"].value for c in "CDEFG"] == [None] * 5
    assert ws["A4"].value == "TC-1"


def test_clear_honours_the_kept_statuses(client, loaded):
    path = paths(client, "Described.xlsx")

    res = client.post(
        "/api/prepare/clear", json={"files": [path], "keep": ["Cancel"], "apply": True}
    )

    assert res.get_json()["results"][0]["plan"]["kept"] == {"Cancel": 1}
    assert load_workbook(path)["Login"]["C5"].value == "対象外"


def test_clear_rejects_a_status_key_the_taxonomy_does_not_have(client, loaded):
    path = paths(client, "Described.xlsx")

    res = client.post("/api/prepare/clear", json={"files": [path], "keep": ["Nope"]})

    assert res.status_code == 400
    assert "Nope" in res.get_json()["error"]


def test_clear_reports_a_workbook_with_no_tool_data_as_an_error(client, loaded):
    path = paths(client, "Bare.xlsx")

    result = client.post("/api/prepare/clear", json={"files": [path]}).get_json()["results"][0]

    assert "TOOL_DATA" in result["error"]


# --- what may be touched ---------------------------------------------------

def test_a_path_outside_the_loaded_source_is_refused(client, loaded, tmp_path):
    outside = write_workbook(tmp_path.parent / "Elsewhere.xlsx", None, {"Login": bare_cells()})

    res = client.post("/api/prepare/tool-data", json={"files": [outside], "apply": True})

    assert res.status_code == 400
    assert has_tool_data(outside) is False


def test_clear_refuses_a_path_outside_the_loaded_source(client, loaded, tmp_path):
    outside = write_workbook(
        tmp_path.parent / "Elsewhere2.xlsx",
        [config_row("Login", "iPhone", 4, 4, cols="A B C D E F G")],
        {"Login": {(4, "A"): "TC-1", (4, "C"): "OK"}},
    )

    res = client.post("/api/prepare/clear", json={"files": [outside], "apply": True})

    assert res.status_code == 400
    assert load_workbook(outside)["Login"]["C4"].value == "OK"


def test_an_empty_file_list_is_an_error(client, loaded):
    res = client.post("/api/prepare/clear", json={"files": []})

    assert res.status_code == 400
