import os

import pytest
from openpyxl import Workbook
from openpyxl.utils import column_index_from_string

from parser.excel_reader import TOOL_DATA_SHEET, TOOL_DATA_COLUMNS

DEFAULT_COLS = "A B C D E F G"

_COL_KEYS = [
    "test_no_col", "scope_col", "result_col", "test_date_col",
    "pic_col", "ticket_id_col", "note_col",
]


def config_row(sheet, device, start_row, end_row, cols=DEFAULT_COLS, **overrides):
    """One TOOL_DATA row. `cols` is the 7 column letters, in TOOL_DATA order."""
    row = {
        "sheet": sheet,
        "device": device,
        "start_row": start_row,
        "end_row": end_row,
        **dict(zip(_COL_KEYS, cols.split())),
    }
    row.update(overrides)
    return row


def write_workbook(path, tool_data, cells):
    """Build an .xlsx. `cells` maps sheet name -> {(row, "A"): value}.

    `tool_data=None` builds a workbook with no TOOL_DATA sheet at all — the
    state a real test case workbook arrives in before anyone has described its
    layout, and the one `prepare.tool_data` exists to fix.
    """
    wb = Workbook()
    default = wb.active
    if tool_data is None:
        wb.remove(default)
    else:
        td = default
        td.title = TOOL_DATA_SHEET
        td.append(TOOL_DATA_COLUMNS)
        for row in tool_data:
            td.append([row.get(name) for name in TOOL_DATA_COLUMNS])

    for sheet_name, sheet_cells in cells.items():
        ws = wb.create_sheet(sheet_name)
        for (row_num, letter), value in sheet_cells.items():
            ws.cell(row=row_num, column=column_index_from_string(letter), value=value)

    wb.save(path)
    return str(path)


@pytest.fixture
def make_workbook(tmp_path):
    def _make(name="TC.xlsx", tool_data=(), cells=None):
        return write_workbook(tmp_path / name, list(tool_data), cells or {})
    return _make


# --- editable configs ------------------------------------------------------
# `config_store.save` writes a file *and* mutates the live singletons, so a
# test that saves anything has to be given somewhere to write and be put back
# afterwards. Both fixtures live here because the store's own tests and the
# endpoint's tests need exactly the same pair.

@pytest.fixture
def restore_configs():
    """Put the live config singletons back after a test has written over them."""
    import config_store
    from tcm.domain.scope import SCOPES
    from tcm.domain.sheet_labels import LABELS
    from tcm.domain.status import STATUS
    from tcm.infrastructure.report.layout import LAYOUT

    from tcm.domain.device import DEVICES

    saved = [(obj, dict(obj.__dict__)) for obj in (STATUS, SCOPES, DEVICES, LABELS, LAYOUT)]
    yield config_store
    for obj, state in saved:
        obj.__dict__.clear()
        obj.__dict__.update(state)


@pytest.fixture
def config_paths(tmp_path, monkeypatch):
    """Point every editable config at a copy in `tmp_path`.

    The repo's own JSON must never be the thing under test: a save that went to
    the wrong place would rewrite the shipped taxonomy.
    """
    import tcm.settings as app_config
    import config_store

    paths = {}
    for name, spec in config_store.CONFIGS.items():
        source = getattr(app_config, spec.setting)
        # Copied under the real file's name: the store labels its refusals with
        # the basename, so a copy called something else would not be testable.
        target = tmp_path / os.path.basename(source)
        with open(source, encoding="utf-8") as f:
            target.write_text(f.read(), encoding="utf-8")
        monkeypatch.setattr(app_config, spec.setting, str(target))
        paths[name] = target
    return paths


#: The scope groups the tests below are written against.
#:
#: Deliberately *not* `parser/scope_groups.json`. That file is editable from the
#: Config view at runtime, which is a feature — so a test asserting that
#: "FPT (JM Support)" classifies as FPT was really asserting that nobody had
#: exercised the feature yet, and broke the moment somebody did. Tests that care
#: what the groups *are* pin them here; tests that care about the shipped file
#: assert its invariants (exactly one fallback, last, never counted) rather than
#: its contents.
FIXED_SCOPES = {
    "groups": [
        {"key": "FPT", "label": "FPT", "match": ["FPT", "FPT (JM Support)"]},
        {"key": "JP", "label": "JP", "match": ["JP"], "excluded": True},
    ],
    "fallback": {"key": "Other", "label": "Other"},
}


@pytest.fixture
def fixed_scopes():
    """Adopt {@link FIXED_SCOPES} for one test, then put the real ones back.

    `adopt` copies onto the live singleton rather than rebinding the name, which
    is what makes this reach the eight modules that imported `SCOPES` by name —
    the same mechanism the Config view's save uses.
    """
    from tcm.domain.scope import SCOPES, ScopeSet

    saved = dict(SCOPES.__dict__)
    SCOPES.adopt(ScopeSet.from_dict(FIXED_SCOPES))
    yield SCOPES
    SCOPES.__dict__.clear()
    SCOPES.__dict__.update(saved)
